"""Create two disposable decks for the native-window smoke test."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

destination = Path(__file__).resolve().parents[1] / "output/presentation-test"
destination.mkdir(parents=True, exist_ok=True)
for ratio, width in [("16-9", 13.333333), ("4-3", 10)]:
    deck = Presentation()
    deck.slide_width = Inches(width)
    deck.slide_height = Inches(7.5)
    for index in range(1, 4):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F3F7FC")
        for text, top, size, color in [
            ("THỬ TRÌNH CHIẾU " + ratio.replace("-", ":"), 1.1, 34, "123C65"),
            (f"SLIDE {index} / 3", 2.6, 72, "197DA3"),
            ("Dùng pointer: trái / phải hoặc Page Up / Page Down", 4.8, 21, "31536F"),
            ("Kiểm tra căn khung, phóng slide và trở về bố cục", 5.5, 18, "526B84"),
        ]:
            box = slide.shapes.add_textbox(Inches(.7), Inches(top), Inches(width - 1.4), Inches(1.4))
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = text
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = size >= 34
            paragraph.font.color.rgb = RGBColor.from_string(color)
    deck.save(destination / f"reading-view-{ratio}.pptx")
print(destination)
